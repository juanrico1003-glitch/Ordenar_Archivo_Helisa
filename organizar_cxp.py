import sys
import os
import glob
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

try:
    import xlrd
except ImportError:
    xlrd = None

# Mapeo de abreviaturas de meses en español
MESES_MAP = {
    'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6,
    'JUL': 7, 'AGO': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12
}
MESES_REV = {
    1: 'ENE', 2: 'FEB', 3: 'MAR', 4: 'ABR', 5: 'MAY', 6: 'JUN',
    7: 'JUL', 8: 'AGO', 9: 'SEP', 10: 'OCT', 11: 'NOV', 12: 'DIC'
}


def parse_fecha(dato):
    """Convierte texto de fecha (ej. '24/AGO/2026' o '24/08/2026') en objeto date de Python."""
    if dato is None:
        return None
    if isinstance(dato, datetime.datetime):
        return dato.date()
    if isinstance(dato, datetime.date):
        return dato
    texto = str(dato).strip().upper()
    if not texto:
        return None
    partes = texto.split('/')
    if len(partes) == 3:
        try:
            d = int(partes[0])
            m_str = partes[1]
            y = int(partes[2])
            if m_str in MESES_MAP:
                return datetime.date(y, MESES_MAP[m_str], d)
            elif m_str.isdigit():
                return datetime.date(y, int(m_str), d)
        except Exception:
            return None
    return None


CATEGORIAS = ['EXTERIOR', 'SERVICIOS PUBLICOS', 'NACIONALES']


def es_documento_excluido(doc_str):
    """Determina si un comprobante debe excluirse de las hojas organizadas.
    
    Se excluyen únicamente:
    - Comprobantes que inician por 'CC', 'ANT' o '02'
    - Comprobantes que inician por '04 LEG' (ej. '04  LEGAGO10', '04  0LEGAGO4')
    
    Todas las demás facturas que inician por '04' (ej. '04  0000031F', '04  FCAR...')
    sí se organizan normalmente.
    """
    if not doc_str:
        return False
    doc = str(doc_str).strip().upper()
    # 1. Comprobantes CC, ANT o 02
    if doc.startswith(('CC', 'ANT', '02')):
        return True
    # 2. Comprobantes 04 LEG únicamente (sin afectar las demás facturas 04)
    doc_sin_espacios = doc.replace(' ', '')
    if doc_sin_espacios.startswith(('04LEG', '040LEG')):
        return True
    return False


def clasificar_hoja(cta):
    """Clasifica la cuenta contable (CTA) en una de las 3 categorías definidas."""
    if not cta:
        return 'NACIONALES'
    cta_str = str(cta).strip()
    if '220510' in cta_str:
        return 'EXTERIOR'
    elif '23355001' in cta_str:
        return 'SERVICIOS PUBLICOS'
    else:
        return 'NACIONALES'


def generar_hoja_cxp(wb, nombre_hoja, filas_registros):
    """Crea y da formato a una hoja organizada (EXTERIOR, SERVICIOS PUBLICOS o NACIONALES)."""
    if nombre_hoja in wb.sheetnames:
        wb.remove(wb[nombre_hoja])
    ws = wb.create_sheet(title=nombre_hoja)

    # 1. Detectar quincenas activas en los registros de esta hoja
    active_quincenas_set = set()
    for (_, _, _, _, fecha_vence, _) in filas_registros:
        if fecha_vence:
            q = 1 if fecha_vence.day <= 15 else 2
            active_quincenas_set.add((fecha_vence.year, fecha_vence.month, q))

    active_quincenas = sorted(active_quincenas_set)

    # 2. Construir dinámicamente los encabezados con su año correspondiente
    encabezados = [
        'CTA', 'PROVEEDOR', 'No. FC', 'FECHA HELISA', 'FECHA', 'F.VENCE',
        'valor', 'TIEMPO CREDITO'
    ]
    for y, m, q in active_quincenas:
        encabezados.append(f"{MESES_REV[m]} Q{q} {y}")
    encabezados.extend(['pago', 'VALOR'])
    total_cols = len(encabezados)

    header_font = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
    for col_idx, h_text in enumerate(encabezados, start=1):
        c = ws.cell(5, col_idx, h_text)
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    # 3. Escribir datos, fórmulas y formatos a partir de la fila 6
    fila_salida = 6
    data_font = Font(name='Aptos Narrow', size=11, bold=False, color='000000')
    currency_fmt = '"$" #,##0.00;-"$" #,##0.00'
    date_fmt = 'dd/mmm/yyyy'

    for idx, (cta, supp, doc, fecha_doc, fecha_vence, valor) in enumerate(filas_registros):
        r_out = fila_salida + idx

        # Col A: CTA
        ws.cell(r_out, 1).value = cta

        # Col B: PROVEEDOR
        ws.cell(r_out, 2).value = supp

        # Col C: No. FC
        if doc:
            c_c = ws.cell(r_out, 3)
            c_c.value = doc
            c_c.alignment = Alignment(horizontal='center')
        else:
            ws.cell(r_out, 3).value = None

        # Col D: FECHA HELISA (=IF(OR(F6="",H6=""),"",F6-H6))
        c_d = ws.cell(r_out, 4)
        c_d.value = f'=IF(OR(F{r_out}="",H{r_out}=""),"",F{r_out}-H{r_out})'
        c_d.number_format = date_fmt
        c_d.alignment = Alignment(horizontal='center')

        # Col E: FECHA
        if fecha_doc:
            c_e = ws.cell(r_out, 5)
            c_e.value = fecha_doc
            c_e.number_format = date_fmt
            c_e.alignment = Alignment(horizontal='center')
        else:
            ws.cell(r_out, 5).value = None

        # Col F: F.VENCE
        if fecha_vence:
            c_f = ws.cell(r_out, 6)
            c_f.value = fecha_vence
            c_f.number_format = date_fmt
            c_f.alignment = Alignment(horizontal='center')
        else:
            ws.cell(r_out, 6).value = None

        # Col G: valor
        c_g = ws.cell(r_out, 7)
        c_g.value = valor
        c_g.number_format = currency_fmt

        # Col H: TIEMPO CREDITO (=IF(OR(E6="",F6=""),"",F6-E6))
        c_h = ws.cell(r_out, 8)
        c_h.value = f'=IF(OR(E{r_out}="",F{r_out}=""),"",F{r_out}-E{r_out})'
        c_h.number_format = '0'
        c_h.alignment = Alignment(horizontal='center')

        # Quincenas dinámicas filtrando por Año, Mes y Quincena
        for q_idx, (y, m, q) in enumerate(active_quincenas):
            col_idx = 9 + q_idx
            cq = ws.cell(r_out, col_idx)
            if q == 1:
                cq.value = f'=IF(AND($F{r_out}<>"",YEAR($F{r_out})={y},MONTH($F{r_out})={m},DAY($F{r_out})<=15),$G{r_out},"")'
            else:
                cq.value = f'=IF(AND($F{r_out}<>"",YEAR($F{r_out})={y},MONTH($F{r_out})={m},DAY($F{r_out})>=16),$G{r_out},"")'
            cq.number_format = currency_fmt

        # Columnas pago y VALOR (vacías)
        col_pago = 9 + len(active_quincenas)
        col_valor = col_pago + 1
        ws.cell(r_out, col_pago).value = None
        ws.cell(r_out, col_valor).value = None

        for c_idx in range(1, total_cols + 1):
            ws.cell(r_out, c_idx).font = data_font

    last_row = fila_salida + len(filas_registros) - 1 if filas_registros else 5

    # 4. Subtotales de cada quincena activa en la Fila 4
    if filas_registros:
        sum_font = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
        for q_idx in range(len(active_quincenas)):
            col_idx = 9 + q_idx
            col_let = get_column_letter(col_idx)
            c_sum = ws.cell(4, col_idx)
            c_sum.value = f'=SUM({col_let}6:{col_let}{last_row})'
            c_sum.font = sum_font
            c_sum.number_format = currency_fmt

    # 5. Configuración estética: Anchos de columna, Autofiltro y Paneles
    col_widths = {
        'A': 24.8, 'B': 55.8, 'C': 18.8, 'D': 14.8, 'E': 13.0,
        'F': 13.0, 'G': 16.8, 'H': 17.8
    }
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w

    for q_idx in range(len(active_quincenas)):
        ws.column_dimensions[get_column_letter(9 + q_idx)].width = 14.5

    col_pago = 9 + len(active_quincenas)
    col_valor = col_pago + 1
    ws.column_dimensions[get_column_letter(col_pago)].width = 14.8
    ws.column_dimensions[get_column_letter(col_valor)].width = 13.0

    last_col_let = get_column_letter(total_cols)
    ws.auto_filter.ref = f'A5:{last_col_let}{max(last_row, 5)}'
    ws.freeze_panes = 'A6'

    return ws


def procesar_archivo_cxp(filepath):
    print(f"\n=======================================================")
    print(f" Procesando: {os.path.basename(filepath)}")
    print(f"=======================================================")

    is_xls = filepath.lower().endswith(".xls")
    is_xlsm = filepath.lower().endswith(".xlsm")

    if is_xls:
        if xlrd is None:
            print(f"ERROR: La librería 'xlrd' no está instalada para leer archivos .xls.")
            print("Por favor ejecuta 'instalar_requisitos.bat'.")
            return False

        try:
            wb_xls = xlrd.open_workbook(filepath)
        except Exception as e:
            print(f"ERROR al leer el archivo binario .xls: {e}")
            return False

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for sname in wb_xls.sheet_names():
            ws_src = wb_xls.sheet_by_name(sname)
            sheet_title = "BASE" if sname == "Sheet 1" else sname
            ws_dst = wb.create_sheet(title=sheet_title)
            for r in range(ws_src.nrows):
                for c in range(ws_src.ncols):
                    ctype = ws_src.cell_type(r, c)
                    val = ws_src.cell_value(r, c)
                    if ctype == xlrd.XL_CELL_DATE:
                        try:
                            val = xlrd.xldate_as_datetime(val, wb_xls.datemode)
                        except Exception:
                            pass
                    elif ctype == xlrd.XL_CELL_EMPTY:
                        val = None
                    if val is not None and val != "":
                        ws_dst.cell(row=r + 1, column=c + 1, value=val)
    else:
        wb = openpyxl.load_workbook(filepath, keep_vba=is_xlsm)

    # 1. Buscar la hoja de origen
    raw_sheet_name = None
    hojas_generadas = set(['Hoja1', 'CUENTAS POR PAGAR'] + CATEGORIAS)

    # Prioridad 1: hoja llamada "BASE"
    for name in wb.sheetnames:
        if name in hojas_generadas:
            continue
        if name.upper() == "BASE":
            raw_sheet_name = name
            break

    # Prioridad 2: nombres conocidos de reportes Helisa
    if not raw_sheet_name:
        for name in wb.sheetnames:
            if name in hojas_generadas:
                continue
            n_up = name.upper()
            if "DETALLADO POR EDADES" in n_up or "DETALLADO CUENTAS CONTABLES" in n_up:
                raw_sheet_name = name
                break

    # Prioridad 3: coincidencia parcial con "DETALLADO", "CUENTAS" o "EDADES"
    if not raw_sheet_name:
        for name in wb.sheetnames:
            if name in hojas_generadas:
                continue
            n_up = name.upper()
            if any(k in n_up for k in ["DETALLADO", "CUENTAS", "EDADES"]):
                raw_sheet_name = name
                break

    # Prioridad 4: inspeccionar encabezados en las primeras filas
    if not raw_sheet_name:
        for name in wb.sheetnames:
            if name in hojas_generadas:
                continue
            ws = wb[name]
            found = False
            for r in range(1, min(25, ws.max_row + 1)):
                for c in range(1, min(16, ws.max_column + 1)):
                    cell_val = str(ws.cell(r, c).value or "")
                    if any(header in cell_val for header in ["FECHA", "Vr. Total", "F.VENCE", "(Dir:"]):
                        raw_sheet_name = name
                        found = True
                        break
                if found:
                    break
            if raw_sheet_name:
                break

    # Prioridad 5: Primera hoja que no sea una hoja generada
    if not raw_sheet_name:
        non_gen = [n for n in wb.sheetnames if n not in hojas_generadas]
        if non_gen:
            raw_sheet_name = non_gen[0]

    if not raw_sheet_name:
        print(f"ERROR: No se encontró la hoja con el reporte crudo en {filepath}")
        return False

    raw_sheet = wb[raw_sheet_name]
    if raw_sheet_name != "BASE":
        if "BASE" in wb.sheetnames:
            wb.remove(wb["BASE"])
        raw_sheet.title = "BASE"

    # 2. Extraer registros del reporte crudo
    curr_supp = None
    curr_cta = None
    extracted_rows = []
    suppliers_set = set()
    total_valor = 0.0
    omitidos_count = 0

    for r in range(1, raw_sheet.max_row + 1):
        val_a = raw_sheet.cell(r, 1).value
        val_f = raw_sheet.cell(r, 6).value   # FECHA
        val_j = raw_sheet.cell(r, 10).value  # Vr. Total
        val_o = raw_sheet.cell(r, 15).value  # F.VENCE

        # Identificar encabezado de Proveedor o Cuenta en columna A
        if val_a and isinstance(val_a, str):
            val_a_str = val_a.strip()
            if any(h in val_a_str for h in ["SODECA LATAM", "Estado detallado", "CUENTA "]):
                continue
            
            if "(Dir:" in val_a_str:
                curr_supp = val_a_str
                curr_cta = None
                suppliers_set.add(curr_supp)
                continue
                
            tokens = val_a_str.split()
            if tokens and tokens[0].isdigit() and len(tokens[0]) >= 4:
                curr_cta = val_a_str
                continue

        # Identificar registros de transacciones Col J
        if val_j is not None and isinstance(val_j, (int, float)):
            doc_str = str(val_a).strip() if val_a is not None else None

            # Omitir comprobantes que inicien por CC, ANT, 02 o 04 LEG
            if es_documento_excluido(doc_str):
                omitidos_count += 1
                continue

            fecha_doc = parse_fecha(val_f)
            fecha_vence = parse_fecha(val_o)
            valor = float(val_j)
            
            extracted_rows.append((curr_cta, curr_supp, doc_str, fecha_doc, fecha_vence, valor))
            total_valor += valor

    if not extracted_rows:
        print("No se encontraron registros para organizar.")
        return False

    # 3. Agrupar registros por categoría contable
    rows_by_category = {cat: [] for cat in CATEGORIAS}
    for row in extracted_rows:
        cta = row[0]
        cat = clasificar_hoja(cta)
        rows_by_category[cat].append(row)

    # 4. Generar cada una de las 3 hojas organizadas
    for cat in CATEGORIAS:
        generar_hoja_cxp(wb, cat, rows_by_category[cat])

    # 5. Limpieza: Eliminar hojas antiguas si existen
    for old_sheet in ["Hoja1", "CUENTAS POR PAGAR"]:
        if old_sheet in wb.sheetnames:
            wb.remove(wb[old_sheet])

    # 6. Activar la primera hoja organizada
    if CATEGORIAS[0] in wb.sheetnames:
        wb.active = wb[CATEGORIAS[0]]

    # 7. Guardar el archivo
    try:
        if is_xls:
            target_path = os.path.splitext(filepath)[0] + ".xlsx"
            wb.save(target_path)
            print(f" -> Archivo organizado guardado como: {os.path.basename(target_path)}")
            try:
                os.remove(filepath)
                print(f" -> Archivo original .xls ({os.path.basename(filepath)}) eliminado para evitar duplicados.")
            except Exception as e:
                print(f" -> Aviso: No se pudo eliminar el archivo .xls original: {e}")
        else:
            wb.save(filepath)
            print(f" -> Archivo guardado exitosamente: {os.path.basename(filepath)}")
    except PermissionError:
        print(f"\n[AVISO] No se pudo guardar '{os.path.basename(filepath)}' porque está abierto en Excel.")
        print("        Por favor cierra el archivo en Excel y vuelve a ejecutar el proceso.\n")
        return False
    except Exception as e:
        print(f"\n[ERROR] Ocurrió un problema al guardar '{os.path.basename(filepath)}': {e}\n")
        return False

    # 8. Reporte en consola
    print(f" -> Proveedores procesados: {len(suppliers_set)}")
    print(f" -> Total registros:        {len(extracted_rows)}")
    if omitidos_count > 0:
        print(f" -> Omitidos (CC/ANT/02/04 LEG): {omitidos_count} (conservados solo en hoja BASE)")
    print(f" -> Suma total de valores:  ${total_valor:,.2f}")
    print(f" -> Desglose por hojas:")
    for cat in CATEGORIAS:
        filas = rows_by_category[cat]
        val_cat = sum(f[5] for f in filas)
        print(f"    * [{cat}]: {len(filas)} registros | Total: ${val_cat:,.2f}")

    return True


def main():
    # Asegurar que el directorio de trabajo sea la carpeta del script o ejecutable
    if getattr(sys, 'frozen', False):
        script_dir = os.path.dirname(os.path.abspath(sys.executable))
    else:
        script_dir = os.path.dirname(os.path.abspath(__file__))

    if len(sys.argv) > 1:
        target_files = [sys.argv[1]]
    else:
        # Busca archivos Excel .xlsx, .xlsm, .xls en la carpeta del script/ejecutable
        excel_patterns = [
            os.path.join(script_dir, "*.xlsx"),
            os.path.join(script_dir, "*.xlsm"),
            os.path.join(script_dir, "*.xls"),
        ]
        target_files = []
        seen = set()
        for pat in excel_patterns:
            for f in glob.glob(pat):
                norm = os.path.normcase(os.path.abspath(f))
                if norm not in seen:
                    seen.add(norm)
                    target_files.append(f)

        # Excluir temporales de Excel o copias de referencia/respaldo
        target_files = [
            f for f in target_files 
            if not os.path.basename(f).startswith("~$") 
            and not f.lower().endswith("_original.xlsx")
            and not f.lower().endswith("_original.xls")
            and not f.lower().endswith("_original.xlsm")
            and not "copia a seguir" in f.lower()
        ]

    if not target_files:
        print(f"No se encontraron archivos Excel (.xlsx / .xlsm / .xls) para procesar en:\n{script_dir}")
        if getattr(sys, 'frozen', False):
            try:
                input("\nPresione Enter para salir...")
            except Exception:
                pass
        return

    for f in target_files:
        procesar_archivo_cxp(f)

    print("\n¡Proceso terminado exitosamente para todos los archivos!")
    if getattr(sys, 'frozen', False):
        try:
            input("\nPresione Enter para salir...")
        except Exception:
            pass


if __name__ == "__main__":
    main()

